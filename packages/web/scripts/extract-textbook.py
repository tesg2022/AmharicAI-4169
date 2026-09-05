"""Extract the rewritten textbook DOCX into structured JSON.

Faithful extraction — no rewriting of the source wording. Units 1-6 already come
from the original course spec, so only units 7-20 are emitted as new content,
alongside the appendices (pronouns, high-value verbs) and the vowel cue table.
"""

import json
import re
import sys
from docx import Document

SRC = "/home/user/Attachments/AmharicAI_Source_Based_Rewritten_Textbook_100_Page_Rewritten_Edition_n1mDmA.docx"
OUT = "/home/user/amharicai/packages/web/src/api/content/textbook.json"

HEAD_OBJ = "Learning objectives"
HEAD_VOCAB = "Core vocabulary"
HEAD_DIALOGUE = "Model dialogue"
HEAD_PRACTICE = "Practice /"
HEAD_CHECK = "Checkpoint /"

DASH = "—"


def clean(s):
    return re.sub(r"\s+", " ", s).strip()


def split_bullet(line):
    """'• ሰላም — sälam — hello/peace' -> ('ሰላም','sälam','hello/peace')"""
    body = line.lstrip("•").strip()
    parts = [clean(p) for p in body.split(DASH)]
    return parts


def main():
    doc = Document(SRC)
    paras = [clean(p.text) for p in doc.paragraphs if clean(p.text)]

    # ---- units -------------------------------------------------------------
    starts = []
    for i, t in enumerate(paras):
        m = re.match(r"^Unit (\d+):\s*(.+)$", t)
        if m:
            starts.append((i, int(m.group(1)), m.group(2).strip()))

    units = []
    for idx, (i, num, title_am) in enumerate(starts):
        end = starts[idx + 1][0] if idx + 1 < len(starts) else len(paras)
        block = paras[i:end]
        title_en = block[1] if len(block) > 1 else ""

        section = None
        objectives, vocab, dialogue = [], [], []
        practice, checkpoint = [], []
        for line in block[2:]:
            if line.startswith(HEAD_OBJ):
                section = "obj"
                continue
            if line.startswith(HEAD_VOCAB):
                section = "vocab"
                continue
            if line.startswith(HEAD_DIALOGUE):
                section = "dlg"
                continue
            if line.startswith(HEAD_PRACTICE):
                section = "practice"
                continue
            if line.startswith(HEAD_CHECK):
                section = "check"
                continue

            if section == "obj" and line.startswith("•"):
                objectives.append(line.lstrip("•").strip())
            elif section == "vocab" and line.startswith("•"):
                p = split_bullet(line)
                if len(p) >= 3:
                    vocab.append(
                        {"amharic": p[0], "transliteration": p[1], "english": DASH.join(p[2:]).strip()}
                    )
                elif len(p) == 2:
                    vocab.append({"amharic": p[0], "transliteration": "", "english": p[1]})
            elif section == "dlg":
                m = re.match(r"^([A-Zab]|[ሀ-፿ ]+):\s*(.*)$", line)
                if m:
                    rest = m.group(2)
                    am, en = (rest.split(DASH, 1) + [""])[:2]
                    dialogue.append(
                        {"speaker": clean(m.group(1)), "amharic": clean(am), "english": clean(en)}
                    )
                elif DASH in line:
                    am, en = line.split(DASH, 1)
                    dialogue.append({"speaker": "", "amharic": clean(am), "english": clean(en)})
                elif line:
                    dialogue.append({"speaker": "", "amharic": line, "english": ""})
            elif section == "practice":
                practice.append(line)
            elif section == "check":
                checkpoint.append(line)

        units.append(
            {
                "number": num,
                "titleAm": title_am,
                "titleEn": title_en,
                "objectives": objectives,
                "vocabulary": vocab,
                "dialogue": dialogue,
                "practice": " ".join(practice),
                "checkpoint": " ".join(checkpoint),
            }
        )

    # units 1-6 already exist from the original course spec
    new_units = [u for u in units if u["number"] >= 7]

    # ---- tables ------------------------------------------------------------
    def rows(t):
        return [[clean(c.text) for c in r.cells] for r in t.rows]

    tables = doc.tables
    pronouns = [
        {"amharic": r[0], "transliteration": r[1], "english": r[2]}
        for r in rows(tables[3])[1:]
        if r[0]
    ]
    verbs = [
        {"infinitive": r[0], "transliteration": r[1], "past3ms": r[2], "meaning": r[3]}
        for r in rows(tables[4])[1:]
        if r[0]
    ]
    vowel_cues = [{"cue": r[0], "reminder": r[1]} for r in rows(tables[7])[1:] if r[0]]
    practice_plan = [
        {"day": r[0], "task": r[1]} for r in rows(tables[55])[1:] if r[0]
    ]
    performance_tasks = [
        {"area": r[0], "task": r[1]} for r in rows(tables[52])[1:] if r[0]
    ]

    data = {
        "units": new_units,
        "pronouns": pronouns,
        "highValueVerbs": verbs,
        "vowelCues": vowel_cues,
        "practicePlan": practice_plan,
        "performanceTasks": performance_tasks,
    }

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"units {len(new_units)} (7-{new_units[-1]['number']})")
    for u in new_units:
        print(
            f"  u{u['number']:>2} {u['titleEn'][:34]:<34} "
            f"obj={len(u['objectives'])} vocab={len(u['vocabulary'])} dlg={len(u['dialogue'])}"
        )
    print(f"pronouns {len(pronouns)}  verbs {len(verbs)}  vowelCues {len(vowel_cues)}")
    print(f"practicePlan {len(practice_plan)}  performanceTasks {len(performance_tasks)}")


if __name__ == "__main__":
    sys.exit(main())
